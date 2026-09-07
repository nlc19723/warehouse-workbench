#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# dump_one.py - Open workbook via COM (Excel/WPS), write ONE .xlsx
# with ALL worksheets (values only, no formatting). Drag .xls/.xlsx/.xlsm.
import sys, os, datetime, re
import win32com.client as win32


# Patterns for date-time strings that COM sometimes returns as text.
_DATE_PATTERNS = [
    re.compile(r'^(\d{4})[-/](\d{1,2})[-/](\d{1,2})\s+(\d{1,2}):(\d{2}):(\d{2})(?:\.\d+)?$'),
    re.compile(r'^(\d{4})年(\d{1,2})月(\d{1,2})日\s*(\d{1,2}):(\d{2}):(\d{2})(?:\.\d+)?$'),
]


def _try_parse_datetime_string(s):
    if not isinstance(s, str):
        return s
    s = s.strip()
    for pat in _DATE_PATTERNS:
        m = pat.match(s)
        if m:
            y, mo, d, h, mi, sec = map(int, m.groups()[:6])
            try:
                dt = datetime.datetime(y, mo, d, h, mi, sec)
                if dt.time() == datetime.time(0, 0, 0):
                    return dt.date()
                return dt
            except ValueError:
                pass
    return s


def _cell_to_value(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, datetime.datetime):
        if v.tzinfo is not None:
            v = v.replace(tzinfo=None)
        if v.time() == datetime.time(0, 0, 0):
            return v.date()
        return v
    if isinstance(v, datetime.time) and v.tzinfo is not None:
        return v.replace(tzinfo=None)
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, str):
        parsed = _try_parse_datetime_string(v)
        if parsed is not v:
            return parsed
    return v


def _rows(raw):
    if raw is None:
        return []
    if not isinstance(raw, tuple):
        return [[raw]]
    if len(raw) > 0 and not isinstance(raw[0], tuple):
        return [list(raw)]
    return [list(r) for r in raw]


def _safe_sheet_name(name, used):
    s = (name or "Sheet")[:31]
    if s in used:
        used[s] += 1
        s = s[:27] + "_" + str(used[s])
    else:
        used[s] = 1
    return s


def launch():
    for pid in ("Excel.Application", "KET.Application", "ET.Application"):
        try:
            app = win32.Dispatch(pid)
            app.Visible = False
            app.DisplayAlerts = False
            try:
                app.ScreenUpdating = False
                app.Calculation = -4135
                app.EnableEvents = False
            except Exception:
                pass
            return app, pid
        except Exception:
            continue
    raise RuntimeError("no Excel/WPS COM available")


def main():
    if len(sys.argv) < 2:
        print("Usage: python dump_one.py <file>")
        sys.exit(2)
    src = os.path.abspath(sys.argv[1])
    if not os.path.exists(src):
        print("[!] file not found: " + src)
        sys.exit(2)
    out_dir = os.path.join(os.path.dirname(src), "dump_out")
    os.makedirs(out_dir, exist_ok=True)
    base = os.path.splitext(os.path.basename(src))[0]
    out_path = os.path.join(out_dir, base + "_extracted.xlsx")

    app = None
    try:
        app, pid = launch()
        print("[*] opened with " + pid)
        wb = app.Workbooks.Open(src, ReadOnly=True)
        try:
            all_sheets = []
            for ws in wb.Worksheets:
                rows = _rows(ws.UsedRange.Value)
                all_sheets.append((ws.Name, rows))
                print("[*] read sheet '" + str(ws.Name) + "': " + str(len(rows)) + " rows")
        finally:
            wb.Close(SaveChanges=False)
    except Exception as e:
        print("[!] open/read failed: " + str(e))
        sys.exit(1)
    finally:
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass

    try:
        from openpyxl import Workbook
    except ImportError:
        print("[!] openpyxl not installed: pip install openpyxl")
        sys.exit(1)

    nwb = Workbook()
    nwb.remove(nwb.active)
    used = {}
    for name, rows in all_sheets:
        ws = nwb.create_sheet(title=_safe_sheet_name(name, used))
        for i, row in enumerate(rows, start=1):
            for j, val in enumerate(row, start=1):
                v = _cell_to_value(val)
                if v is not None:
                    ws.cell(row=i, column=j, value=v)
    nwb.save(out_path)
    sz = os.path.getsize(out_path)
    print("[OK] saved: " + out_path)
    print("[OK] " + str(len(all_sheets)) + " sheets, " + str(sz) + " bytes")


if __name__ == "__main__":
    main()
